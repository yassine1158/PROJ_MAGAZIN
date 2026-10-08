"""Exports Excel."""
from django.http import HttpResponse
from django.utils import timezone
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter


def reponse_excel(titre, entetes, lignes, nom_fichier):
    wb = Workbook()
    ws = wb.active
    ws.title = titre[:31]
    ws.append(entetes)
    for cellule in ws[1]:
        cellule.font = Font(bold=True, color='FFFFFF')
        cellule.fill = PatternFill('solid', fgColor='1F497D')
    for ligne in lignes:
        ws.append([float(v) if hasattr(v, 'as_tuple') else v for v in ligne])
    for i, entete in enumerate(entetes, start=1):
        largeur = max([len(str(entete))] + [len(str(c.value or '')) for c in ws[get_column_letter(i)]])
        ws.column_dimensions[get_column_letter(i)].width = min(largeur + 2, 50)
    ws.freeze_panes = 'A2'

    response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    date = timezone.localdate().strftime('%Y-%m-%d')
    response['Content-Disposition'] = f'attachment; filename="{nom_fichier}_{date}.xlsx"'
    wb.save(response)
    return response


def etat_stock(articles):
    return reponse_excel(
        'État du stock',
        ['Code', 'Désignation', 'Catégorie', 'Unité', 'Emplacement', 'Stock', 'Stock min',
         'Prix moyen', 'Valeur', 'Alerte'],
        [[a.code, a.designation, a.categorie.nom, a.unite, a.emplacement, a.stock, a.stock_min,
          a.prix_moyen, a.valeur_stock, 'OUI' if a.en_alerte else ''] for a in articles],
        'etat_stock',
    )


def mouvements(qs):
    return reponse_excel(
        'Mouvements',
        ['Date', 'Bon', 'Type', 'Code', 'Désignation', 'Quantité', 'Unité', 'Stock après',
         'Prix unitaire', 'Chantier', 'Engin', 'Utilisateur'],
        [[timezone.localtime(m.date).strftime('%d/%m/%Y %H:%M'), m.reference, m.get_type_display(),
          m.article.code, m.article.designation, m.quantite, m.article.unite, m.stock_apres,
          m.prix_unitaire, m.chantier.nom if m.chantier else '', m.engin.code if m.engin else '',
          str(m.utilisateur or '')] for m in qs.select_related('article', 'chantier', 'engin', 'utilisateur')],
        'mouvements',
    )
